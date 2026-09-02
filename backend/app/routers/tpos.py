"""Training & Placement Officer (TPO) directory — for campus hiring outreach + hiring vendors."""
from __future__ import annotations

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import models, schemas
from ..database import get_db
from ..services.recruitment import log

router = APIRouter(prefix="/api/tpos", tags=["tpos"])


@router.get("", response_model=list[schemas.TPOOut])
def list_tpos(db: Session = Depends(get_db)):
    return db.scalars(select(models.TPO).order_by(models.TPO.college, models.TPO.name)).all()


# ── Bulk import (Excel .xlsx / .csv) — vendors & placement cells from a partner spreadsheet ──────
_TPO_MAX_BYTES = 5 * 1024 * 1024
# Import columns. The importer maps these to core fields; ANY other column is kept under `details`
# so the partner tables can show the rich fields (industries, TAT, fee, programs, CTC, …).
TPO_IMPORT_COLUMNS = [
    "Category", "Organization Name", "Point of Contact", "Designation", "Email", "Phone",
    "Website/LinkedIn", "Country", "City", "Geographies / Programs", "Industries / Disciplines",
    "Hiring Types / Season", "TAT / Internship Duration", "Fee / Min Stipend", "Min CTC", "Notes",
]
_TPO_CORE = {"category", "organization name", "point of contact", "designation", "email", "phone",
             "website/linkedin", "country", "city", "notes"}


def _parse_sheet(filename: str, content: bytes) -> list[dict]:
    import io
    if (filename or "").lower().endswith(".csv"):
        import csv
        rows = [list(r) for r in csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")))]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
    if not rows:
        return []
    headers = [str(h or "").strip().lower() for h in rows[0]]
    return [{headers[i]: ("" if (i >= len(r) or r[i] is None) else str(r[i]).strip())
             for i in range(len(headers)) if headers[i]} for r in rows[1:]]


def _tpo_kind(category: str) -> str:
    c = (category or "").lower()
    return "vendor" if ("vendor" in c or "agency" in c or "recruit" in c) else "college"


@router.get("/import/template")
def tpo_import_template():
    """Download the .xlsx template for importing vendors + placement cells."""
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Partners"
    ws.append(TPO_IMPORT_COLUMNS)
    ws.append(["Vendor", "Acme Recruiters", "Jane Doe", "Talent Head", "jane@acme.com", "+91 98000 00000",
               "www.acme.com", "India", "Bengaluru", "India, Middle East", "IT / Technology",
               "Full-Time, Contractual", "24-48 hrs", "8.33% of CTC", "", "90-day replacement"])
    ws.append(["Placement cell", "IIM Example", "John Roe", "Placement Officer", "tpo@iim.edu", "+91 99000 00000",
               "www.iim.edu", "India", "Indore", "PGP, MBA", "Management / Business Studies",
               "Aug-Oct 2025", "2 months", "₹60,000 stipend", "15 LPA", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="partners_import_template.xlsx"'})


@router.post("/import")
def import_tpos(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Bulk-create vendors / placement cells from an .xlsx/.csv. Core columns map to fields; every
    other column is stored under `details` so it shows in the partner table. Dedup by email."""
    content = file.file.read(_TPO_MAX_BYTES + 1)
    if len(content) > _TPO_MAX_BYTES:
        raise HTTPException(400, "File too large (max 5 MB).")
    try:
        rows = _parse_sheet(file.filename or "", content)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"Couldn't read the spreadsheet ({str(exc)[:80]}). Use the template.")
    if not rows:
        raise HTTPException(400, "No data rows found.")
    created = skipped = 0
    for d in rows:
        org = (d.get("organization name") or d.get("organization") or "").strip()
        poc = (d.get("point of contact") or d.get("poc") or "").strip()
        email = (d.get("email") or "").strip().lower()
        if not org and not poc:
            skipped += 1
            continue
        if email and db.scalar(select(models.TPO).where(models.TPO.email == email)):
            skipped += 1
            continue
        addr = ", ".join(p for p in [(d.get("city") or "").strip(), (d.get("country") or "").strip()] if p)
        details = {k.title(): v for k, v in d.items() if k not in _TPO_CORE and v}
        db.add(models.TPO(
            kind=_tpo_kind(d.get("category")),
            name=poc, college=org, email=email,
            phone=(d.get("phone") or "").strip(),
            linkedin=(d.get("website/linkedin") or d.get("website") or "").strip(),
            designation=(d.get("designation") or "").strip(),
            address=addr, notes=(d.get("notes") or "").strip(), details=details,
        ))
        created += 1
    log(db, "tpos.imported", "tpo", None, {"created": created, "skipped": skipped})
    db.commit()
    return {"created": created, "skipped": skipped, "total_rows": len(rows)}


@router.post("/intake", status_code=201)
def partner_intake(payload: schemas.TPOIntake, db: Session = Depends(get_db)):
    """PUBLIC (no auth) — a hiring vendor / college fills the public form and the row lands
    directly in the partner directory. Allow-listed in main.py's auth gate. Deduped by email so
    a double-submit doesn't create two rows."""
    org = (payload.organization or "").strip()
    poc = (payload.contact_name or "").strip()
    email = (payload.email or "").strip().lower()
    if not org:
        raise HTTPException(422, "Organization name is required.")
    if not poc:
        raise HTTPException(422, "A point-of-contact name is required.")
    if email and db.scalar(select(models.TPO).where(models.TPO.email == email)):
        # Idempotent: a repeat submission from the same email just succeeds silently.
        return {"ok": True, "duplicate": True}
    addr = ", ".join(p for p in [(payload.city or "").strip(), (payload.country or "").strip()] if p)
    details = {k: v for k, v in (payload.details or {}).items() if k and str(v).strip()}
    tpo = models.TPO(
        kind=_tpo_kind(payload.category),
        name=poc[:160], college=org[:200], email=email[:200],
        phone=(payload.phone or "").strip()[:60],
        linkedin=(payload.website or "").strip()[:300],
        designation=(payload.designation or "").strip()[:160],
        address=addr, notes=(payload.notes or "").strip(), details=details,
    )
    db.add(tpo)
    db.flush()
    log(db, "tpo.intake", "tpo", tpo.id, {"kind": tpo.kind, "college": tpo.college})
    db.commit()
    return {"ok": True, "id": tpo.id}


@router.post("", response_model=schemas.TPOOut, status_code=201)
def create_tpo(payload: schemas.TPOCreate, db: Session = Depends(get_db)):
    if not (payload.name or "").strip() and not (payload.college or "").strip():
        raise HTTPException(422, "A name or college is required.")
    tpo = models.TPO(
        kind=(payload.kind or "college").strip().lower() or "college",
        name=payload.name.strip(),
        college=payload.college.strip(),
        email=(payload.email or "").strip(),
        phone=(payload.phone or "").strip(),
        linkedin=(payload.linkedin or "").strip(),
        designation=(payload.designation or "").strip(),
        address=(payload.address or "").strip(),
        notes=(payload.notes or "").strip(),
    )
    db.add(tpo)
    db.flush()
    log(db, "tpo.created", "tpo", tpo.id, {"college": tpo.college})
    db.commit()
    db.refresh(tpo)
    return tpo


@router.patch("/{tpo_id}", response_model=schemas.TPOOut)
def update_tpo(tpo_id: int, payload: schemas.TPOUpdate, db: Session = Depends(get_db)):
    tpo = db.get(models.TPO, tpo_id)
    if not tpo:
        raise HTTPException(404, "TPO not found")
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(tpo, k, (v or "").strip() if isinstance(v, str) else v)
    db.commit()
    db.refresh(tpo)
    return tpo


@router.delete("/{tpo_id}", status_code=204)
def delete_tpo(tpo_id: int, db: Session = Depends(get_db)):
    tpo = db.get(models.TPO, tpo_id)
    if not tpo:
        raise HTTPException(404, "TPO not found")
    db.delete(tpo)
    log(db, "tpo.deleted", "tpo", tpo_id, {})
    db.commit()
